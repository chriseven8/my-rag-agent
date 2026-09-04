from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.ingestion.loader import load_text
from app.ingestion.splitter import split_sections

FIXTURE_PDF = Path(__file__).resolve().parent / "fixtures" / "minimal.pdf"


def _make_docx(path: Path) -> Path:
    """造一个带 1/2 级标题 + 正文 + 表格的 docx。"""
    doc = Document()
    doc.add_heading("产品介绍", level=1)
    doc.add_paragraph("苹果 价格 是 3000 元")
    doc.add_heading("规格", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "价格"
    doc.save(path)
    return path


def test_docx_heading_and_body_and_table(tmp_path):
    path = _make_docx(tmp_path / "产品.docx")
    text = load_text(str(path))
    assert "# 产品介绍" in text
    assert "苹果 价格 是 3000 元" in text
    assert "## 规格" in text
    # 表格按顺序穿插,行内单元格以 | 连接
    assert "名称 | 价格" in text


def test_docx_table_section_path_from_heading(tmp_path):
    path = _make_docx(tmp_path / "产品.docx")
    text = load_text(str(path))
    sections = split_sections(text)
    hit = next(s for s in sections if "名称 | 价格" in s[1])
    # 表格应归属到它前面最近的标题层级:产品介绍 > 规格
    assert hit[0] == "产品介绍 > 规格"


def test_docx_plain_paragraph_without_heading(tmp_path):
    doc = Document()
    doc.add_paragraph("只有一段普通文本 123")
    path = tmp_path / "plain.docx"
    doc.save(path)
    text = load_text(str(path))
    assert "只有一段普通文本 123" in text
    assert not text.lstrip().startswith("#")


def _make_xlsx(path: Path) -> Path:
    """造一个带数据 sheet「库存」和一个空 sheet「空表」的 xlsx。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "库存"
    ws.append(["名称", "价格"])
    ws.append(["苹果", 3000])
    ws.append(["香蕉", None, None, None])  # 行尾空列应被裁掉
    wb.create_sheet("空表")  # 无任何单元格
    wb.save(path)
    return path


def test_xlsx_sheet_heading_values_and_trailing_trim(tmp_path):
    path = _make_xlsx(tmp_path / "价目.xlsx")
    text = load_text(str(path))
    assert "# 库存" in text
    assert "名称 | 价格" in text
    assert "苹果 | 3000" in text
    # 行尾空列被裁掉,不产出 "香蕉 |"
    assert "香蕉 |" not in text
    assert "空表" not in text  # 空 sheet 不产出标题行


def test_xlsx_rows_section_path_is_sheet_name(tmp_path):
    path = _make_xlsx(tmp_path / "价目.xlsx")
    text = load_text(str(path))
    sections = split_sections(text)
    hit = next(s for s in sections if "苹果 | 3000" in s[1])
    assert hit[0] == "库存"


def test_pdf_fixture_extracts_text():
    assert load_text(str(FIXTURE_PDF)) == "Hello RAG PDF"
